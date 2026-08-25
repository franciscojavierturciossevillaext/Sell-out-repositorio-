"""
actualizador_cartera.py
=======================
Extrae la tabla Excel llamada "DATA" desde el archivo
  Reporte Cartera Colombia 2026.xlsx
y actualiza (o previsualiza) el archivo CSV de destino
  DATA2026.csv

Uso
---
  # Modo DEMO (no sobrescribe nada):
  python actualizador_cartera.py --demo

  # Modo REAL (reemplaza el CSV después de confirmar la demo):
  python actualizador_cartera.py --ejecutar

Configuración
-------------
  Ajusta las rutas en la sección CONFIGURACIÓN más abajo si cambian
  las rutas de OneDrive o el nombre del archivo.
"""

from __future__ import annotations

import argparse
import csv
import logging
import os
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd

# ---------------------------------------------------------------------------
# CONFIGURACIÓN – ajusta estas rutas si cambian
# ---------------------------------------------------------------------------
ONEDRIVE_BASE = Path(os.environ.get(
    "ONEDRIVE_BASE",
    r"C:\Users\EADKD\OneDrive - Bayer\CARMEN PACA\tablas Dashboard cartera",
))

EXCEL_ORIGEN = ONEDRIVE_BASE / "Reporte Cartera Colombia 2026.xlsx"
CSV_DESTINO  = ONEDRIVE_BASE / "DATA2026.csv"

NOMBRE_TABLA = "DATA"          # nombre exacto de la tabla en el Excel
FILAS_PREVIEW = 5              # filas a mostrar en la demo
LOG_FILE = Path(__file__).parent / "actualizador_cartera.log"
# ---------------------------------------------------------------------------

# ---------------------------------------------------------------------------
# Logging
# ---------------------------------------------------------------------------
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(levelname)-8s  %(message)s",
    handlers=[
        logging.FileHandler(LOG_FILE, encoding="utf-8"),
        logging.StreamHandler(sys.stdout),
    ],
)
logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Funciones de utilidad
# ---------------------------------------------------------------------------

def _detectar_encoding_csv(ruta: Path) -> str:
    """Intenta detectar la codificación del CSV existente."""
    try:
        import chardet  # opcional; si no está instalado, usa utf-8
        with open(ruta, "rb") as f:
            resultado = chardet.detect(f.read(50_000))
        enc = resultado.get("encoding") or "utf-8"
        logger.info("Codificación detectada del CSV: %s", enc)
        return enc
    except ImportError:
        logger.info("chardet no disponible; usando utf-8 como codificación del CSV.")
        return "utf-8"


def _detectar_delimitador_csv(ruta: Path, encoding: str) -> str:
    """Detecta el delimitador del CSV existente."""
    with open(ruta, "r", encoding=encoding, errors="replace") as f:
        muestra = f.read(4096)
    sniffer = csv.Sniffer()
    try:
        dialecto = sniffer.sniff(muestra)
        logger.info("Delimitador detectado: %r", dialecto.delimiter)
        return dialecto.delimiter
    except csv.Error:
        logger.warning("No se pudo detectar el delimitador; usando ','.")
        return ","


def _verificar_archivo_abierto(ruta: Path) -> None:
    """
    En Windows, intentar renombrar temporalmente detecta si el archivo está
    bloqueado por otra aplicación (Excel, OneDrive sync, etc.).
    En otros SO el bloqueo no aplica, así que se omite silenciosamente.
    """
    if not ruta.exists():
        return
    if sys.platform != "win32":
        return
    tmp = ruta.with_suffix(".tmp_lock_check")
    try:
        ruta.rename(tmp)
        tmp.rename(ruta)
    except PermissionError:
        raise PermissionError(
            f"El archivo '{ruta.name}' está abierto o siendo sincronizado por OneDrive. "
            "Ciérralo y vuelve a ejecutar el script."
        )


def leer_tabla_excel(ruta_excel: Path, nombre_tabla: str) -> pd.DataFrame:
    """
    Lee la tabla nombrada *nombre_tabla* desde el Excel.
    Lanza ValueError si la tabla no existe.
    """
    logger.info("Abriendo: %s", ruta_excel)
    _verificar_archivo_abierto(ruta_excel)

    wb = openpyxl.load_workbook(ruta_excel, read_only=False, data_only=True)

    tabla_ref = None
    hoja_tabla = None

    for hoja in wb.worksheets:
        for tabla in hoja.tables.values():
            t_nombre = tabla.displayName if hasattr(tabla, "displayName") else str(tabla)
            if t_nombre.upper() == nombre_tabla.upper():
                tabla_ref = tabla
                hoja_tabla = hoja
                break
        if tabla_ref:
            break

    if tabla_ref is None:
        nombres_encontrados = [
            (t.displayName if hasattr(t, "displayName") else str(t))
            for ws in wb.worksheets for t in ws.tables.values()
        ]
        raise ValueError(
            f"No se encontró una tabla llamada '{nombre_tabla}' en el Excel.\n"
            f"Tablas encontradas: {nombres_encontrados or '(ninguna)'}"
        )

    # Leer el rango de la tabla
    filas = list(hoja_tabla[tabla_ref.ref])
    encabezados = [celda.value for celda in filas[0]]
    datos = [[celda.value for celda in fila] for fila in filas[1:]]

    logger.info(
        "Tabla '%s' encontrada en hoja '%s' | rango: %s | %d filas x %d columnas",
        nombre_tabla, hoja_tabla.title, tabla_ref.ref, len(datos), len(encabezados),
    )

    df = pd.DataFrame(datos, columns=encabezados)
    return df


def _info_csv_actual(ruta: Path) -> dict:
    """Devuelve metadatos del CSV actual para comparación."""
    if not ruta.exists():
        return {}
    enc = _detectar_encoding_csv(ruta)
    sep = _detectar_delimitador_csv(ruta, enc)
    df_actual = pd.read_csv(ruta, encoding=enc, sep=sep, nrows=0)
    return {
        "encoding": enc,
        "sep": sep,
        "columnas": list(df_actual.columns),
        "num_columnas": len(df_actual.columns),
    }


# ---------------------------------------------------------------------------
# Modos de ejecución
# ---------------------------------------------------------------------------

def modo_demo(excel: Path, csv_dest: Path, nombre_tabla: str) -> None:
    """
    Lee la tabla, muestra vista previa y compara con el CSV actual.
    NO escribe nada en disco.
    """
    logger.info("=" * 60)
    logger.info("MODO DEMO / PRUEBA  –  ningún archivo será modificado")
    logger.info("=" * 60)

    df_nuevo = leer_tabla_excel(excel, nombre_tabla)

    print("\n" + "=" * 60)
    print(f"  Vista previa – primeras {FILAS_PREVIEW} filas de la tabla '{nombre_tabla}'")
    print("=" * 60)
    print(df_nuevo.head(FILAS_PREVIEW).to_string(index=False))
    print(f"\nTotal: {len(df_nuevo):,} filas  ×  {len(df_nuevo.columns)} columnas")
    print(f"Columnas: {list(df_nuevo.columns)}")

    # Comparar con CSV actual
    info_csv = _info_csv_actual(csv_dest)
    if info_csv:
        print("\n" + "-" * 60)
        print("  Comparación con DATA2026.csv actual")
        print("-" * 60)
        cols_csv  = set(info_csv["columnas"])
        cols_xlsx = set(df_nuevo.columns)
        coinciden = cols_csv == cols_xlsx

        print(f"  Columnas CSV actual : {info_csv['num_columnas']}")
        print(f"  Columnas tabla Excel: {len(df_nuevo.columns)}")
        print(f"  Formato compatible  : {'✅ SÍ' if coinciden else '⚠️  NO (ver diferencias abajo)'}")

        if not coinciden:
            solo_csv  = cols_csv  - cols_xlsx
            solo_xlsx = cols_xlsx - cols_csv
            if solo_csv:
                print(f"  Solo en CSV   : {sorted(solo_csv)}")
            if solo_xlsx:
                print(f"  Solo en Excel : {sorted(solo_xlsx)}")

        print(f"  Codificación CSV    : {info_csv['encoding']}")
        print(f"  Delimitador CSV     : {info_csv['sep']!r}")
    else:
        print(f"\n  ⚠️  El archivo '{csv_dest.name}' no existe aún; se creará en modo real.")

    print("\n  DEMO completada. Para aplicar los cambios ejecuta:")
    print("  python actualizador_cartera.py --ejecutar\n")
    logger.info("DEMO finalizada sin cambios.")


def modo_ejecutar(excel: Path, csv_dest: Path, nombre_tabla: str) -> None:
    """
    Lee la tabla y reemplaza por completo el CSV de destino.
    """
    logger.info("=" * 60)
    logger.info("MODO REAL  –  el CSV será reemplazado")
    logger.info("=" * 60)

    # Detectar formato del CSV existente antes de sobrescribir
    info_csv = _info_csv_actual(csv_dest)
    encoding = info_csv.get("encoding", "utf-8")
    sep      = info_csv.get("sep", ",")

    # Leer tabla
    df_nuevo = leer_tabla_excel(excel, nombre_tabla)

    # Verificar que el CSV de destino no esté abierto
    _verificar_archivo_abierto(csv_dest)

    # Escribir
    df_nuevo.to_csv(
        csv_dest,
        index=False,
        encoding=encoding,
        sep=sep,
    )

    logger.info(
        "CSV actualizado: %s  |  %d filas escritas  |  encoding=%s  |  sep=%r",
        csv_dest, len(df_nuevo), encoding, sep,
    )
    print(f"\n✅  Archivo actualizado: {csv_dest}")
    print(f"   Filas escritas : {len(df_nuevo):,}")
    print(f"   Columnas       : {len(df_nuevo.columns)}")
    print(f"   Codificación   : {encoding}")
    print(f"   Delimitador    : {sep!r}\n")


# ---------------------------------------------------------------------------
# Punto de entrada
# ---------------------------------------------------------------------------

def main() -> None:
    parser = argparse.ArgumentParser(
        description="Actualiza DATA2026.csv desde la tabla 'DATA' del Excel de cartera."
    )
    grupo = parser.add_mutually_exclusive_group(required=True)
    grupo.add_argument(
        "--demo",
        action="store_true",
        help="Muestra vista previa sin modificar ningún archivo.",
    )
    grupo.add_argument(
        "--ejecutar",
        action="store_true",
        help="Reemplaza el CSV con los datos actualizados.",
    )
    parser.add_argument(
        "--excel",
        type=Path,
        default=EXCEL_ORIGEN,
        help="Ruta al Excel de origen (sobreescribe la ruta por defecto).",
    )
    parser.add_argument(
        "--csv",
        type=Path,
        default=CSV_DESTINO,
        help="Ruta al CSV de destino (sobreescribe la ruta por defecto).",
    )
    parser.add_argument(
        "--tabla",
        default=NOMBRE_TABLA,
        help=f"Nombre de la tabla Excel a extraer (default: {NOMBRE_TABLA}).",
    )
    args = parser.parse_args()

    logger.info(
        "Inicio | modo=%s | excel=%s | csv=%s | tabla=%s",
        "DEMO" if args.demo else "REAL",
        args.excel,
        args.csv,
        args.tabla,
    )

    try:
        if not args.excel.exists():
            raise FileNotFoundError(
                f"No se encontró el archivo de origen:\n  {args.excel}\n"
                "Verifica que OneDrive esté sincronizado y la ruta sea correcta."
            )

        if args.demo:
            modo_demo(args.excel, args.csv, args.tabla)
        else:
            modo_ejecutar(args.excel, args.csv, args.tabla)

    except (FileNotFoundError, ValueError, PermissionError) as exc:
        logger.error("ERROR: %s", exc)
        print(f"\n❌  ERROR: {exc}\n")
        sys.exit(1)
    except Exception as exc:
        logger.exception("Error inesperado: %s", exc)
        print(f"\n❌  Error inesperado: {exc}\n")
        sys.exit(1)


if __name__ == "__main__":
    main()
