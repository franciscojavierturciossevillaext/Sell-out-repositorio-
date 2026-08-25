"""
Document parser service.

Supports:
  - Plain text / Markdown  (.txt, .md)  – built-in
  - CSV                    (.csv)       – built-in (pandas)
  - JSON                   (.json)      – built-in
  - PDF                    (.pdf)       – pypdf (optional dep)
  - DOCX                   (.docx)      – python-docx (optional dep)
  - PPTX                   (.pptx)      – python-pptx (optional dep)
  - XLSX                   (.xlsx)      – openpyxl / pandas (optional dep)

Each parser returns a plain string that is then chunked by the caller.
"""

import io
import json
import logging
from pathlib import Path

from app.utils.file_helpers import get_file_category

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Public entry-point
# ---------------------------------------------------------------------------

def parse_document(file_name: str, content: bytes) -> str:
    """
    Parse *content* (raw bytes of a file) and return extracted plain text.
    Raises ``ValueError`` if the file type is not supported.
    """
    category = get_file_category(file_name)
    parsers = {
        "text": _parse_text,
        "csv": _parse_csv,
        "json": _parse_json,
        "pdf": _parse_pdf,
        "docx": _parse_docx,
        "pptx": _parse_pptx,
        "xlsx": _parse_xlsx,
    }
    parser = parsers.get(category)
    if parser is None:
        raise ValueError(f"Unsupported file type for '{file_name}' (category='{category}').")
    return parser(content)


# ---------------------------------------------------------------------------
# Individual parsers
# ---------------------------------------------------------------------------

def _parse_text(content: bytes) -> str:
    return content.decode("utf-8", errors="replace")


def _parse_json(content: bytes) -> str:
    try:
        data = json.loads(content)
        return json.dumps(data, ensure_ascii=False, indent=2)
    except json.JSONDecodeError:
        return content.decode("utf-8", errors="replace")


def _parse_csv(content: bytes) -> str:
    try:
        import pandas as pd  # type: ignore
        df = pd.read_csv(io.BytesIO(content))
        return df.to_string(index=False)
    except ImportError:
        logger.warning("pandas not installed – returning raw CSV text.")
        return content.decode("utf-8", errors="replace")
    except Exception as exc:
        logger.error("CSV parse error: %s", exc)
        return content.decode("utf-8", errors="replace")


def _parse_pdf(content: bytes) -> str:
    try:
        from pypdf import PdfReader  # type: ignore
        reader = PdfReader(io.BytesIO(content))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n\n".join(pages)
    except ImportError:
        logger.warning("pypdf not installed – PDF parsing not available.")
        return "[PDF content – install pypdf to enable parsing]"
    except Exception as exc:
        logger.error("PDF parse error: %s", exc)
        return f"[PDF parse error: {exc}]"


def _parse_docx(content: bytes) -> str:
    try:
        from docx import Document  # type: ignore
        doc = Document(io.BytesIO(content))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n\n".join(paragraphs)
    except ImportError:
        logger.warning("python-docx not installed – DOCX parsing not available.")
        return "[DOCX content – install python-docx to enable parsing]"
    except Exception as exc:
        logger.error("DOCX parse error: %s", exc)
        return f"[DOCX parse error: {exc}]"


def _parse_pptx(content: bytes) -> str:
    try:
        from pptx import Presentation  # type: ignore
        prs = Presentation(io.BytesIO(content))
        slides_text: list[str] = []
        for i, slide in enumerate(prs.slides, start=1):
            texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text)
            if texts:
                slides_text.append(f"--- Slide {i} ---\n" + "\n".join(texts))
        return "\n\n".join(slides_text)
    except ImportError:
        logger.warning("python-pptx not installed – PPTX parsing not available.")
        return "[PPTX content – install python-pptx to enable parsing]"
    except Exception as exc:
        logger.error("PPTX parse error: %s", exc)
        return f"[PPTX parse error: {exc}]"


def _parse_xlsx(content: bytes) -> str:
    try:
        import pandas as pd  # type: ignore
        xl = pd.ExcelFile(io.BytesIO(content))
        parts: list[str] = []
        for sheet in xl.sheet_names:
            df = xl.parse(sheet)
            parts.append(f"=== Sheet: {sheet} ===\n{df.to_string(index=False)}")
        return "\n\n".join(parts)
    except ImportError:
        logger.warning("pandas/openpyxl not installed – XLSX parsing not available.")
        return "[XLSX content – install pandas + openpyxl to enable parsing]"
    except Exception as exc:
        logger.error("XLSX parse error: %s", exc)
        return f"[XLSX parse error: {exc}]"
